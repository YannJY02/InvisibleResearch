#!/usr/bin/env python3
"""Exploratory exact-ISSN comparison: OpenAlex 2025-08 versus Scopus 2026-08.

Requires openpyxl. Reads a completed BigQuery tabledata acquisition, preserves
every journal row, and retains all exact candidates rather than selecting a title.
Generated files belong in this pipeline's ignored artifacts directory.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import date, datetime, timezone
import gzip
import hashlib
import json
from pathlib import Path
import re
import urllib.request

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parents[1]
REFERENCE_URL = (
    "https://downloads.ctfassets.net/o78em1y1w4i4/7xtaTxNiNcWRTeZkV86eNy/"
    "69cf2d506c905dc299531fdc93049dbb/ext_list_Aug_2026.xlsx"
)
LANDING_URL = "https://www.elsevier.com/products/scopus/content"
REFERENCE_NAME = "ext_list_Aug_2026.xlsx"
REFERENCE_SHA256 = "11e81f686401c89fbef28de31d1880388bb7122f35c89e09db6e1848237e9afb"
MAIN_SHEET = "Scopus Sources Aug. 2026"
DISCONTINUED_SHEET = "Discontinued Titles Aug. 2026"
CONFERENCE_SHEET = "Serial Conf. Proc. with Profile"
ACCEPTED_SHEET = "Accepted Titles Aug. 2026"


def sha256(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def dumps(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize_issn(value: object) -> tuple[str | None, str]:
    """Return checksum-valid hyphenated ISSN; never repair a bad check digit.

    Excel numeric cells can omit an initial zero, so pad those cells to eight
    characters. String fields must already contain eight ISSN characters after
    removing whitespace and the normal ISSN hyphen.
    """
    if value is None or cell(value).strip() == "":
        return None, "missing"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and not value.is_integer():
            return None, "nonintegral_numeric_cell"
        compact = cell(value).zfill(8)
    else:
        compact = re.sub(r"[\s-]", "", cell(value)).upper()
    if not re.fullmatch(r"[0-9]{7}[0-9X]", compact):
        return None, "invalid_format"
    digits = [int(char) for char in compact[:7]] + [10 if compact[7] == "X" else int(compact[7])]
    if sum(digit * weight for digit, weight in zip(digits, range(8, 0, -1))) % 11:
        return None, "invalid_checksum"
    return compact[:4] + "-" + compact[4:], "valid"


def download_reference(output: Path) -> dict:
    path = output / REFERENCE_NAME
    metadata_path = output / "download.json"
    if not path.exists():
        partial = path.with_suffix(path.suffix + ".partial")
        with urllib.request.urlopen(REFERENCE_URL, timeout=120) as response, partial.open("wb") as stream:
            headers = dict(response.headers.items())
            status = response.status
            while chunk := response.read(1024 * 1024):
                stream.write(chunk)
        expected_length = next((v for k, v in headers.items() if k.lower() == "content-length"), None)
        if expected_length and partial.stat().st_size != int(expected_length):
            raise ValueError("Reference download length does not match Content-Length")
        if sha256(partial) != REFERENCE_SHA256:
            raise ValueError("Downloaded reference differs from the pinned August 2026 artifact")
        partial.replace(path)
        write_json(metadata_path, {
            "source_url": REFERENCE_URL, "landing_url": LANDING_URL,
            "source_title": "Scopus source title list - August 2026",
            "reference_version": "2026-08", "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
            "http_status": status, "http_headers": headers, "bytes": path.stat().st_size,
            "sha256": REFERENCE_SHA256, "local_filename": REFERENCE_NAME,
        })
    if sha256(path) != REFERENCE_SHA256:
        raise ValueError("Local reference differs from the pinned August 2026 artifact")
    if not metadata_path.exists():
        raise ValueError("The reference exists but its HTTP provenance metadata is absent")
    metadata = json.loads(metadata_path.read_text())
    if metadata["sha256"] != REFERENCE_SHA256 or metadata["source_url"] != REFERENCE_URL:
        raise ValueError("Reference provenance does not match the pinned artifact")
    return metadata


def load_reference(path: Path) -> tuple[list[dict], list[dict], dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    inventory = {"worksheets": {s.title: {"max_row": s.max_row, "max_column": s.max_column}
                               for s in workbook.worksheets}, "parsed_sheets": {}}
    records, accepted = [], []
    for sheet_name, header_row in [(MAIN_SHEET, 1), (DISCONTINUED_SHEET, 2),
                                   (CONFERENCE_SHEET, 1), (ACCEPTED_SHEET, 3)]:
        sheet = workbook[sheet_name]
        rows, ids, statuses, types, validation, numeric_cells = 0, [], Counter(), Counter(), Counter(), 0
        headers = [cell(v) for v in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
        for number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not any(v is not None for v in values):
                continue
            rows += 1
            pending = sheet_name == ACCEPTED_SHEET
            id_value, title, issn, eissn = (("", values[0], values[1], values[2]) if pending
                                           else (values[0], values[1], values[2], values[3]))
            record = {
                "sheet": sheet_name, "excel_row": number, "source_id": cell(id_value),
                "title": cell(title), "raw_issn": cell(issn), "raw_eissn": cell(eissn),
                "issn_cell_type": type(issn).__name__, "eissn_cell_type": type(eissn).__name__,
                "valid_issns": [], "invalid_issns": [], "active_or_inactive": "",
                "source_type": "", "source_type_origin": "not_stated_in_sheet", "coverage": "",
                "discontinued_marker": "", "indexation_change": "", "publisher": "",
            }
            for field, value in [("ISSN", issn), ("EISSN", eissn)]:
                normalized, reason = normalize_issn(value)
                validation[reason] += 1
                numeric_cells += isinstance(value, (int, float)) and not isinstance(value, bool)
                if normalized:
                    record["valid_issns"].append(normalized)
                elif reason != "missing":
                    record["invalid_issns"].append({"field": field, "raw": cell(value), "reason": reason})
            record["valid_issns"] = sorted(set(record["valid_issns"]))
            if sheet_name == MAIN_SHEET:
                record.update(active_or_inactive=cell(values[4]), coverage=cell(values[5]),
                              discontinued_marker=cell(values[6]), source_type=cell(values[12]),
                              source_type_origin="Source Type column", publisher=cell(values[18]),
                              title_history=cell(values[13]), related_titles=[cell(v) for v in values[14:18] if v])
            elif sheet_name == DISCONTINUED_SHEET:
                record.update(publisher=cell(values[4]), indexation_change=cell(values[5]),
                              final_coverage={headers[i]: cell(values[i]) for i in range(6, 10)})
            elif sheet_name == CONFERENCE_SHEET:
                record.update(source_type="Serial Conference Proceedings", source_type_origin="worksheet identity",
                              discontinued_marker=cell(values[4]), coverage=cell(values[5]))
            else:
                record.update(acceptance_date=cell(values[3]), publisher=cell(values[4]))
            if record["source_id"]:
                ids.append(record["source_id"])
            statuses[record["active_or_inactive"] or "not_stated"] += 1
            types[record["source_type"] or "not_stated"] += 1
            (accepted if pending else records).append(record)
        inventory["parsed_sheets"][sheet_name] = {
            "header_row": header_row, "headers": headers, "nonempty_records": rows,
            "unique_source_ids": len(set(ids)), "missing_source_ids": rows - len(ids),
            "duplicate_source_ids": [key for key, count in Counter(ids).items() if count > 1],
            "active_or_inactive": dict(statuses), "source_types": dict(types),
            "issn_cell_validation": dict(validation), "numeric_issn_cells": numeric_cells,
        }
    main_ids = {r["source_id"] for r in records if r["sheet"] == MAIN_SHEET}
    inventory["discontinued_source_ids_absent_from_main"] = len({
        r["source_id"] for r in records if r["sheet"] == DISCONTINUED_SHEET and r["source_id"]} - main_ids)
    inventory["scope_note"] = (
        "The workbook's More Info. Medline sheet says the Source List omits some historical serials: "
        "pre-1996 inactive titles except parent-child relationships; post-1995 titles with fewer than "
        "10 articles awaiting a profile; titles with scattered or incomplete content; and standalone "
        "books, conferences and reports. Absence is not proof of never being indexed. Accepted titles "
        "are pending and are not included in the candidate reference. The separate all-conference "
        "proceedings worksheet has no ISSNs and is not parsed for this journal comparison."
    )
    workbook.close()
    return records, accepted, inventory


def oa_issns(row: dict) -> tuple[list[str], list[dict]]:
    # The downloaded table stores issn as a JSON-array string, observed directly.
    raw_values = json.loads(row["issn"]) if row["issn"] else []
    if not isinstance(raw_values, list) or any(not isinstance(v, str) for v in raw_values):
        raise ValueError(f"Unexpected OpenAlex issn representation for source {row['id']}")
    values = [("issn", value) for value in raw_values]
    if row["issn_l"]:
        values.append(("issn_l", row["issn_l"]))
    valid, invalid = set(), []
    for field, value in values:
        normalized, reason = normalize_issn(value)
        if normalized:
            valid.add(normalized)
        elif reason != "missing":
            invalid.append({"field": field, "raw": value, "reason": reason})
    return sorted(valid), invalid


def build_candidates(valid_issns: list[str], index: dict, records_by_key: dict) -> list[dict]:
    matched = defaultdict(set)
    for issn in valid_issns:
        for key in index.get(issn, []):
            matched[key].add(issn)
    candidates = []
    for key in sorted(matched):
        records = records_by_key[key]
        states = sorted({r["active_or_inactive"] for r in records if r["active_or_inactive"]})
        source_types = sorted({r["source_type"] for r in records if r["source_type"]})
        discontinued = any(r["sheet"] == DISCONTINUED_SHEET or r["discontinued_marker"] for r in records)
        conflict = len(states) > 1 or ("Active" in states and discontinued)
        state = "status_conflict" if conflict else states[0].lower() if states else "status_unknown"
        candidates.append({
            "candidate_key": key, "source_id": records[0]["source_id"],
            "matched_issns": sorted(matched[key]), "candidate_status": state,
            "active_or_inactive_values": states, "source_types": source_types,
            "has_discontinued_evidence": discontinued, "status_conflict": conflict,
            "reference_rows": records,
        })
    return candidates


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=BASE / "artifacts/bigquery-tabledata-2025-08/journals.csv.gz")
    parser.add_argument("--output", type=Path, default=BASE / "artifacts/scopus-2026-08")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    provenance = download_reference(args.output)
    records, accepted, inventory = load_reference(args.output / REFERENCE_NAME)
    write_json(args.output / "reference-inventory.json", inventory)
    for name, data in [("reference-rows.jsonl.gz", records), ("accepted-pending-rows.jsonl.gz", accepted)]:
        with gzip.open(args.output / name, "wt") as stream:
            for record in data:
                stream.write(dumps(record) + "\n")
    manifest_path = args.input.parent / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    input_sha256 = sha256(args.input)
    if input_sha256 != manifest["sha256"][args.input.name]:
        raise ValueError("Input hash differs from completed acquisition manifest")
    index, grouped, pending_index = defaultdict(set), defaultdict(list), defaultdict(list)
    for record in records:
        key = "id:" + record["source_id"] if record["source_id"] else f"row:{record['sheet']}:{record['excel_row']}"
        grouped[key].append(record)
        for issn in record["valid_issns"]:
            index[issn].add(key)
    for record in accepted:
        for issn in record["valid_issns"]:
            pending_index[issn].append(record)
    output_path = args.output / "openalex-journals-scopus-candidates.csv.gz"
    partial = output_path.with_suffix(output_path.suffix + ".partial")
    extra_fields = ["valid_issns_json", "invalid_issns_json", "scopus_match_status", "scopus_journal_match_status", "scopus_candidate_count",
                    "scopus_candidate_ids_json", "scopus_matched_issns_json", "scopus_candidates_json",
                    "scopus_candidate_status_counts_json", "scopus_source_types_json",
                    "scopus_journal_candidate_count", "scopus_has_discontinued_evidence", "scopus_has_status_conflict",
                    "scopus_has_candidate_without_source_id", "scopus_accepted_pending_json"]
    counts, metrics, candidate_types, candidate_states = Counter(), Counter(), Counter(), Counter()
    journal_counts, candidate_usage = Counter(), Counter()
    examples, seen_ids, input_ids_hash = defaultdict(list), set(), hashlib.sha256()
    with gzip.open(args.input, "rt", newline="") as source, gzip.open(partial, "wt", newline="") as destination:
        reader = csv.DictReader(source)
        original_fields = reader.fieldnames
        if set(original_fields) & set(extra_fields):
            raise ValueError("Input already contains Scopus enrichment fields")
        writer = csv.DictWriter(destination, fieldnames=original_fields + extra_fields)
        writer.writeheader()
        for row in reader:
            if row["type"] != "journal" or not row["id"] or row["id"] in seen_ids:
                raise ValueError("Input must contain unique, nonempty journal IDs only")
            seen_ids.add(row["id"])
            input_ids_hash.update((row["id"] + "\n").encode())
            valid, invalid = oa_issns(row)
            candidates = build_candidates(valid, index, grouped)
            status = ("no_valid_issn" if not valid else "not_listed_in_reference" if not candidates
                      else "matched_multiple" if len(candidates) > 1 else "matched_" + candidates[0]["candidate_status"])
            journal_candidates = [c for c in candidates if "Journal" in c["source_types"]]
            journal_status = ("no_valid_issn" if not valid else "no_journal_candidate_in_reference" if not journal_candidates
                              else "matched_multiple" if len(journal_candidates) > 1
                              else "matched_" + journal_candidates[0]["candidate_status"])
            journal_counts[journal_status] += 1
            candidate_usage.update(c["candidate_key"] for c in candidates)
            pending = {f"{r['sheet']}:{r['excel_row']}": r for issn in valid for r in pending_index.get(issn, [])}
            states = Counter(c["candidate_status"] for c in candidates)
            types = sorted({t for c in candidates for t in c["source_types"]})
            discontinued = any(c["has_discontinued_evidence"] for c in candidates)
            conflict = any(c["status_conflict"] for c in candidates)
            counts[status] += 1
            metrics["rows"] += 1
            metrics["rows_with_valid_issn"] += bool(valid)
            metrics["rows_with_invalid_issn"] += bool(invalid)
            metrics["rows_with_any_candidate"] += bool(candidates)
            metrics["rows_with_any_journal_candidate"] += bool(journal_candidates)
            metrics["rows_with_pending_acceptance_match"] += bool(pending)
            metrics["not_listed_with_pending_acceptance_match"] += bool(pending) and status == "not_listed_in_reference"
            metrics["rows_with_discontinued_evidence"] += discontinued
            metrics["rows_with_status_conflict"] += conflict
            metrics["rows_with_candidate_missing_source_id"] += any(not c["source_id"] for c in candidates)
            metrics["rows_with_any_active_candidate"] += any(c["candidate_status"] == "active" for c in candidates)
            metrics["rows_with_active_journal_candidate"] += any(c["candidate_status"] == "active" and "Journal" in c["source_types"] for c in candidates)
            metrics["rows_with_only_explicit_nonjournal_candidates"] += bool(candidates) and all(c["source_types"] and "Journal" not in c["source_types"] for c in candidates)
            candidate_types.update(types)
            candidate_states.update(states)
            enriched = {**row, "valid_issns_json": dumps(valid), "invalid_issns_json": dumps(invalid),
                        "scopus_match_status": status, "scopus_journal_match_status": journal_status,
                        "scopus_candidate_count": len(candidates),
                        "scopus_candidate_ids_json": dumps(sorted({c["source_id"] for c in candidates if c["source_id"]})),
                        "scopus_matched_issns_json": dumps(sorted({i for c in candidates for i in c["matched_issns"]})),
                        "scopus_candidates_json": dumps(candidates), "scopus_candidate_status_counts_json": dumps(states),
                        "scopus_source_types_json": dumps(types),
                        "scopus_journal_candidate_count": sum("Journal" in c["source_types"] for c in candidates),
                        "scopus_has_discontinued_evidence": str(discontinued).lower(),
                        "scopus_has_status_conflict": str(conflict).lower(),
                        "scopus_has_candidate_without_source_id": str(any(not c["source_id"] for c in candidates)).lower(),
                        "scopus_accepted_pending_json": dumps(list(pending.values()))}
            writer.writerow(enriched)
            if len(examples[status]) < 3:
                examples[status].append({"openalex_id": row["id"], "title": row["display_name"], "valid_issns": valid,
                                         "candidate_ids": json.loads(enriched["scopus_candidate_ids_json"]),
                                         "candidate_statuses": dict(states), "source_types": types})
    if metrics["rows"] != manifest["journal_rows"]:
        raise ValueError("Journal row count differs from acquisition manifest")
    output_ids_hash, output_rows = hashlib.sha256(), 0
    # Read back and compare every original cell, not only the row count.
    with gzip.open(args.input, "rt", newline="") as source, gzip.open(partial, "rt", newline="") as destination:
        source_reader, result_reader = csv.DictReader(source), csv.DictReader(destination)
        for result in result_reader:
            original = next(source_reader, None)
            if original is None or any(original[k] != result[k] for k in original_fields):
                raise ValueError("An original OpenAlex field changed during enrichment")
            json.loads(result["scopus_candidates_json"])
            output_rows += 1
            output_ids_hash.update((result["id"] + "\n").encode())
        if next(source_reader, None) is not None:
            raise ValueError("An original OpenAlex row is missing from output")
    if output_rows != metrics["rows"] or input_ids_hash.hexdigest() != output_ids_hash.hexdigest():
        raise ValueError("Output IDs or rows do not match the input")
    partial.replace(output_path)
    write_json(args.output / "examples.json", examples)
    summary = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "analysis_status": "Exploratory Analysis; not a same-year predictive outcome",
        "input": {"path": str(args.input), "source": manifest["source"], "snapshot_label": "2025-08",
                  "sha256": input_sha256, "manifest_sha256": sha256(manifest_path), "rows": metrics["rows"]},
        "reference": provenance, "reference_inventory": inventory, "match_status_counts": dict(counts),
        "strict_journal_match_status_counts": dict(journal_counts),
        "strict_journal_definition": "A candidate must have the exact Source Type value Journal in the main sheet. "
                                     "Trade Journal, Book Series, conference profiles and unknown types are excluded. "
                                     "The status counts cover all original OpenAlex rows, not only matches.",
        "metrics": dict(metrics), "rows_by_matched_source_type_nonexclusive": dict(candidate_types),
        "candidate_status_counts_including_repeated_candidates_across_openalex_rows": dict(candidate_states),
        "reference_candidate_keys": len(grouped), "reference_unique_valid_issns": len(index),
        "reference_candidate_keys_matched": len(candidate_usage),
        "reference_candidate_keys_matched_by_multiple_openalex_rows": sum(v > 1 for v in candidate_usage.values()),
        "reference_issns_linked_to_multiple_candidate_keys": sum(len(v) > 1 for v in index.values()),
        "validation": {"output_rows": output_rows, "input_unique_ids": len(seen_ids),
                       "ordered_id_sha256_input": input_ids_hash.hexdigest(),
                       "ordered_id_sha256_output": output_ids_hash.hexdigest(),
                       "all_original_cells_and_row_order_unchanged": True, "input_manifest_hash_verified": True},
        "output": {"path": str(output_path), "sha256": sha256(output_path), "bytes": output_path.stat().st_size},
        "method": "Exact checksum-valid ISSN overlap using OpenAlex issn JSON array union issn_l. No title matching. "
                  "Group Scopus rows by nonempty Source ID; unidentified rows retain their sheet/row key. "
                  "Main, discontinued and serial-conference-profile sheets are candidates; accepted titles are pending only. "
                  "Active/Inactive are copied only from the main sheet; conflicts with discontinued evidence are marked. "
                  "Multiple candidate identities are not resolved automatically.",
        "limitations": [
            "OpenAlex snapshot 2025-08 and Scopus reference 2026-08 differ by one year; not a same-year outcome.",
            "The official reference excludes some indexed historical sources; not_listed_in_reference is not proof of non-indexing.",
            "A valid ISSN and exact overlap do not adjudicate source identity, title history, or classification differences.",
            "ISSN-L is used as an additional identifier, not as a mapping of all media editions from the ISSN Registry.",
            "Unknown reference status/type remain unknown; pending acceptance does not prove inclusion.",
            "The publicly downloadable workbook is retained locally as provenance and is not republished.",
        ],
    }
    write_json(args.output / "summary.json", summary)
    lines = ["# Scopus ISSN candidate comparison — Exploratory Analysis", "",
             f"Compared **{metrics['rows']:,} OpenAlex journal rows (2025-08)** with the official **Scopus August 2026** workbook.",
             "This is an exploratory comparison across different dates, not a same-year predictive outcome.", "",
             f"Official landing page: {LANDING_URL}", f"Official reference: {REFERENCE_URL}",
             f"Reference SHA256: `{REFERENCE_SHA256}`", "",
             "## Results", "", "| Status | OpenAlex rows |", "|---|---:|"]
    lines += [f"| {status} | {count:,} |" for status, count in sorted(counts.items())]
    lines += ["", "### Strict Journal candidates", "", summary["strict_journal_definition"], "",
              "| Status after restricting candidates to Scopus Journal | OpenAlex rows |", "|---|---:|"]
    lines += [f"| {status} | {count:,} |" for status, count in sorted(journal_counts.items())]
    lines += ["", f"Rows with checksum-valid ISSN: {metrics['rows_with_valid_issn']:,} of {metrics['rows']:,}. "
              f"Rows with any candidate: {metrics['rows_with_any_candidate']:,}; "
              f"rows with any strict Journal candidate: {metrics['rows_with_any_journal_candidate']:,}. "
              f"Rows with an explicitly active Journal candidate: {metrics['rows_with_active_journal_candidate']:,}. "
              f"Rows with only explicitly non-Journal candidates: {metrics['rows_with_only_explicit_nonjournal_candidates']:,}.",
              f"Rows with a pending-acceptance match: {metrics['rows_with_pending_acceptance_match']:,}; "
              f"of these, {metrics['not_listed_with_pending_acceptance_match']:,} have no included-source candidate.", "",
              "## Method and interpretation", "",
              summary["method"], "", "Normalize case, whitespace and standard hyphens; validate the ISSN check digit. "
              "Excel numeric ISSN cells are left-padded to eight characters before validation. "
              "Invalid identifiers are retained in diagnostic JSON and never matched. "
              "Every candidate preserves the worksheet, one-based Excel row, raw identifiers, source title, "
              "original status, source type when available, coverage, and discontinued evidence.", "",
              "`matched_multiple` means more than one candidate identity; multiple worksheets with the same Source ID "
              "are one candidate. `matched_status_unknown` means the matching supplement provides no Active/Inactive "
              "value. `matched_status_conflict` means available active and discontinued evidence conflict. "
              "Unknown Source IDs are retained using their sheet and row, without guessing an identity.", ""]
    lines += ["- " + limitation for limitation in summary["limitations"]]
    lines += ["", "The workbook's own scope note: " + inventory["scope_note"], "",
              "## Verification", "", f"Read back all {output_rows:,} output rows and verified every original cell and row order "
              "against the input. IDs remain unique; ordered ID hashes match. Both input and reference hashes were checked. "
              "The acquisition manifest is the completion gate; no partial acquisition file was used for matching.", "",
              "## Examples", "", "The full examples are in `examples.json`; candidate details are embedded as JSON "
              "in `openalex-journals-scopus-candidates.csv.gz`.", ""]
    for status, samples in sorted(examples.items()):
        lines.append(f"- **{status}**: " + "; ".join(
            f"OpenAlex {r['openalex_id']} — {r['title']} (Scopus IDs: {', '.join(r['candidate_ids']) or 'none'})"
            for r in samples))
    (args.output / "findings.md").write_text("\n".join(lines) + "\n")
    print(json.dumps({"output": str(output_path), "match_status_counts": counts, "metrics": metrics,
                      "validation": summary["validation"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
